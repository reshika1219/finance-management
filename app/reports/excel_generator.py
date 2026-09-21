from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.services.monthly_service import get_monthly_totals
from app.services.event_service import get_events_by_month
from app.services.settings_service import get_settings
from app.utils.date_utils import get_month_year_display, format_display_date


def generate_monthly_excel(year: int, month: int, file_path: str, db_path: str = None) -> str:
    """
    Generates a clean, professional monthly Excel (.xlsx) report for the given year and month.
    Includes:
    - Sheet 1: Monthly Summary
    - Sheet 2: Event Details
    - Hidden Sheet: _ReportMetadata
    """
    wb = openpyxl.Workbook()

    totals = get_monthly_totals(year, month, db_path)
    events = get_events_by_month(year, month, db_path)
    settings = get_settings(db_path)
    business_name = settings.business_name.upper()
    month_display = get_month_year_display(year, month)

    # Styles
    title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    subtitle_font = Font(name="Calibri", size=12, bold=True, color="0284C7")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    highlight_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    profit_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    thick_bottom_side = Side(border_style="medium", color="1E293B")
    double_bottom_side = Side(border_style="double", color="1E293B")

    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    summary_border = Border(top=thin_border_side, bottom=thick_bottom_side)
    total_border = Border(top=thin_border_side, bottom=double_bottom_side)

    # ----------------------------------------------------
    # SHEET 1: Monthly Summary
    # ----------------------------------------------------
    ws1 = wb.active
    ws1.title = "Monthly Summary"
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = business_name
    ws1["A1"].font = title_font
    ws1["A2"] = "MONTHLY FINANCIAL SUMMARY"
    ws1["A2"].font = subtitle_font
    ws1["A3"] = month_display
    ws1["A3"].font = bold_font

    ws1.append([])  # Blank row 4

    headers_ws1 = ["Item", "Amount (Rs.)"]
    ws1.append(headers_ws1)
    for col_idx in range(1, 3):
        cell = ws1.cell(row=5, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left" if col_idx == 1 else "right", vertical="center")

    summary_rows = [
        ("Total Income", float(totals.total_income), True, highlight_fill),
        ("Employee Salaries", float(totals.employee_salaries), False, None),
        ("Transportation", float(totals.transportation), False, None),
        ("Food for Employees", float(totals.food), False, None),
        ("Supplier Payments", float(totals.supplier_payments), False, None),
        ("Utilities / Others", float(totals.utilities_others), False, None),
        ("Rent", float(totals.rent), False, None),
        ("Electricity Bill", float(totals.electricity), False, None),
        ("Water Bill", float(totals.water), False, None),
        ("Telephone Bill", float(totals.telephone), False, None),
        ("Monthly Others", float(totals.monthly_others), False, None),
        ("Total Expenditure", float(totals.total_expenditure), True, highlight_fill),
        ("Net Profit", float(totals.net_profit), True, profit_fill),
    ]

    for item_name, val, is_bold, fill in summary_rows:
        row_num = ws1.max_row + 1
        c1 = ws1.cell(row=row_num, column=1, value=item_name)
        c2 = ws1.cell(row=row_num, column=2, value=val)

        c1.font = bold_font if is_bold else normal_font
        c2.font = bold_font if is_bold else normal_font

        c2.number_format = "#,##0.00"
        c2.alignment = Alignment(horizontal="right")

        if fill:
            c1.fill = fill
            c2.fill = fill

        if item_name in ("Total Income", "Total Expenditure"):
            c1.border = summary_border
            c2.border = summary_border
        elif item_name == "Net Profit":
            c1.border = total_border
            c2.border = total_border
        else:
            c1.border = thin_border
            c2.border = thin_border

    ws1.column_dimensions["A"].width = 30
    ws1.column_dimensions["B"].width = 22

    # ----------------------------------------------------
    # SHEET 2: Event Details
    # ----------------------------------------------------
    ws2 = wb.create_sheet(title="Event Details")
    ws2.views.sheetView[0].showGridLines = True

    headers_ws2 = [
        "Date", "Event", "Client", "Location", "Income",
        "Salaries", "Transportation", "Food", "Suppliers",
        "Utilities/Others", "Total Expenditure", "Profit"
    ]
    ws2.append(headers_ws2)
    for col_idx in range(1, len(headers_ws2) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx == 1 else ("right" if col_idx >= 5 else "left"), vertical="center")

    for ev in events:
        r_vals = [
            format_display_date(ev.event_date),
            ev.event_name,
            ev.client_name,
            ev.location or "",
            float(ev.income),
            float(ev.employee_salary),
            float(ev.transportation),
            float(ev.food),
            float(ev.supplier_payments),
            float(ev.utilities_others),
            float(ev.total_expenditure),
            float(ev.profit),
        ]
        ws2.append(r_vals)
        r_idx = ws2.max_row
        for col_idx in range(1, len(headers_ws2) + 1):
            cell = ws2.cell(row=r_idx, column=col_idx)
            cell.font = normal_font
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx >= 5:
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")

    # Totals Row at bottom of Event Details
    if events:
        tot_row_idx = ws2.max_row + 1
        ws2.cell(row=tot_row_idx, column=1, value="TOTAL").font = bold_font
        ws2.cell(row=tot_row_idx, column=1).alignment = Alignment(horizontal="center")

        ws2.cell(row=tot_row_idx, column=5, value=float(totals.total_income)).number_format = "#,##0.00"
        ws2.cell(row=tot_row_idx, column=6, value=float(totals.employee_salaries)).number_format = "#,##0.00"
        ws2.cell(row=tot_row_idx, column=7, value=float(totals.transportation)).number_format = "#,##0.00"
        ws2.cell(row=tot_row_idx, column=8, value=float(totals.food)).number_format = "#,##0.00"
        ws2.cell(row=tot_row_idx, column=9, value=float(totals.supplier_payments)).number_format = "#,##0.00"
        ws2.cell(row=tot_row_idx, column=10, value=float(totals.utilities_others)).number_format = "#,##0.00"
        ws2.cell(row=tot_row_idx, column=11, value=float(totals.total_event_expenditure)).number_format = "#,##0.00"
        ws2.cell(row=tot_row_idx, column=12, value=float(totals.total_income - totals.total_event_expenditure)).number_format = "#,##0.00"

        for col_idx in range(1, len(headers_ws2) + 1):
            cell = ws2.cell(row=tot_row_idx, column=col_idx)
            cell.font = bold_font
            cell.fill = highlight_fill
            cell.border = total_border
            if col_idx >= 5:
                cell.alignment = Alignment(horizontal="right")

    # Auto column width for Sheet 2
    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # ----------------------------------------------------
    # SHEET 3: Hidden Metadata Sheet (_ReportMetadata)
    # ----------------------------------------------------
    ws_meta = wb.create_sheet(title="_ReportMetadata")
    ws_meta.sheet_state = "hidden"

    metadata = [
        ("report_type", "monthly"),
        ("application", "Chirathma Flora"),
        ("format_version", "1"),
        ("year", str(year)),
        ("month", str(month)),
        ("total_income", str(totals.total_income)),
        ("total_expenditure", str(totals.total_expenditure)),
        ("profit", str(totals.net_profit)),
    ]

    for k, v in metadata:
        ws_meta.append([k, v])

    wb.save(file_path)
    return file_path


def generate_combined_excel(report, file_path: str, db_path: str = None) -> str:
    """
    Generates a clean, professional combined financial report (.xlsx) for a range of months.
    Contains month-level totals and grand total summary.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financial Summary"
    ws.views.sheetView[0].showGridLines = True

    settings = get_settings(db_path)
    business_name = settings.business_name.upper()

    # Styles
    title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    subtitle_font = Font(name="Calibri", size=12, bold=True, color="0284C7")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    highlight_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    double_bottom_side = Side(border_style="double", color="1E293B")

    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    total_border = Border(top=thin_border_side, bottom=double_bottom_side)

    ws["A1"] = business_name
    ws["A1"].font = title_font
    ws["A2"] = "FINANCIAL SUMMARY"
    ws["A2"].font = subtitle_font
    ws["A3"] = f"Reporting Period: {report.period_label}"
    ws["A3"].font = bold_font

    ws.append([])  # Blank row 4

    headers = ["Month", "Total Income (Rs.)", "Total Expenditure (Rs.)", "Profit (Rs.)"]
    ws.append(headers)
    for col_idx in range(1, 5):
        cell = ws.cell(row=5, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left" if col_idx == 1 else "right", vertical="center")

    for r in report.rows:
        ws.append([r.month_name, float(r.total_income), float(r.total_expenditure), float(r.profit)])
        r_idx = ws.max_row
        for col_idx in range(1, 5):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = normal_font
            cell.border = thin_border
            if col_idx >= 2:
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")

    # TOTAL ROW
    tot_row_idx = ws.max_row + 1
    ws.cell(row=tot_row_idx, column=1, value="TOTAL").font = bold_font
    ws.cell(row=tot_row_idx, column=2, value=float(report.total_income)).number_format = "#,##0.00"
    ws.cell(row=tot_row_idx, column=3, value=float(report.total_expenditure)).number_format = "#,##0.00"
    ws.cell(row=tot_row_idx, column=4, value=float(report.total_profit)).number_format = "#,##0.00"

    for col_idx in range(1, 5):
        cell = ws.cell(row=tot_row_idx, column=col_idx)
        cell.font = bold_font
        cell.fill = highlight_fill
        cell.border = total_border
        if col_idx >= 2:
            cell.alignment = Alignment(horizontal="right")

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 24

    wb.save(file_path)
    return file_path

