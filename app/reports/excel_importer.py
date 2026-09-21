import openpyxl
from decimal import Decimal
from app.services.report_service import MonthReportRow, CombinedReport
from app.utils.currency import parse_currency
from app.utils.date_utils import get_month_year_display


def read_monthly_excel_file(file_path: str) -> MonthReportRow:
    """
    Reads an exported Chirathma Flora monthly report Excel file.
    Validates _ReportMetadata sheet and returns MonthReportRow.
    Raises ValueError if invalid or unsupported.
    """
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        raise ValueError(f"Unable to open file: {str(e)}")

    if "_ReportMetadata" not in wb.sheetnames:
        raise ValueError(
            "Unable to Import File\n\nThis does not appear to be a valid Chirathma Flora monthly report."
        )

    ws_meta = wb["_ReportMetadata"]
    meta_dict = {}
    for row in ws_meta.iter_rows(values_only=True):
        if row and len(row) >= 2 and row[0] is not None:
            meta_dict[str(row[0]).strip()] = str(row[1]).strip() if row[1] is not None else ""

    app_name = meta_dict.get("application", "")
    rep_type = meta_dict.get("report_type", "")
    version = meta_dict.get("format_version", "1")

    if app_name != "Chirathma Flora" or rep_type != "monthly":
        raise ValueError(
            "Unable to Import File\n\nThis does not appear to be a valid Chirathma Flora monthly report."
        )

    try:
        year = int(meta_dict.get("year", "0"))
        month = int(meta_dict.get("month", "0"))
        income = parse_currency(meta_dict.get("total_income", "0.00"))
        expenditure = parse_currency(meta_dict.get("total_expenditure", "0.00"))
        profit = parse_currency(meta_dict.get("profit", "0.00"))
    except Exception:
        raise ValueError("Invalid numerical metadata values in monthly report file.")

    if year <= 0 or month < 1 or month > 12:
        raise ValueError("Invalid month or year metadata in monthly report file.")

    month_name = get_month_year_display(year, month)
    return MonthReportRow(
        year=year,
        month=month,
        month_name=month_name,
        total_income=income,
        total_expenditure=expenditure,
        profit=profit,
    )


def build_combined_report_from_files(file_paths: list[str]) -> CombinedReport:
    """
    Imports multiple monthly Excel files, validates metadata, checks for duplicate months,
    sorts chronologically, and returns a CombinedReport object.
    """
    if not file_paths:
        raise ValueError("No files selected.")

    rows: list[MonthReportRow] = []
    seen_months: dict[tuple[int, int], str] = {}

    for path in file_paths:
        row = read_monthly_excel_file(path)
        key = (row.year, row.month)
        if key in seen_months:
            raise ValueError(
                f"Duplicate Month\n\n{row.month_name} has been selected more than once.\n\nPlease remove the duplicate file and try again."
            )
        seen_months[key] = path
        rows.append(row)

    # Sort chronologically by year and month
    rows.sort(key=lambda r: (r.year, r.month))

    grand_income = Decimal("0.00")
    grand_expenditure = Decimal("0.00")
    grand_profit = Decimal("0.00")

    for r in rows:
        grand_income += r.total_income
        grand_expenditure += r.total_expenditure
        grand_profit += r.profit

    first_month = rows[0].month_name
    last_month = rows[-1].month_name
    if len(rows) == 1:
        period_label = first_month
    else:
        period_label = f"{first_month} – {last_month}"

    return CombinedReport(
        period_label=period_label,
        rows=rows,
        total_income=grand_income,
        total_expenditure=grand_expenditure,
        total_profit=grand_profit,
    )
